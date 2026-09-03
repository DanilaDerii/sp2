"""Slide-aware PPTX text extraction worker."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from ..common.models import ExtractedDocument, ExtractedPage


def _normalized_text(value: str) -> str:
    return " ".join(value.split()).strip()


def _shape_position(shape: Any) -> tuple[int, int]:
    return int(getattr(shape, "top", 0)), int(getattr(shape, "left", 0))


def _pptx_shape_text_blocks(shape: Any) -> list[str]:
    blocks: list[str] = []

    if getattr(shape, "has_text_frame", False):
        text = _normalized_text(shape.text)
        if text:
            blocks.append(text)

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            row_text = " | ".join(
                text
                for cell in row.cells
                if (text := _normalized_text(cell.text))
            )
            if row_text:
                blocks.append(row_text)

    child_shapes = getattr(shape, "shapes", None)
    if child_shapes is not None:
        for child in sorted(child_shapes, key=_shape_position):
            blocks.extend(_pptx_shape_text_blocks(child))

    return blocks


def _extract_pptx_pages(source_path: Path) -> list[ExtractedPage]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "PPTX extraction requires python-pptx from environment/requirements.txt"
        ) from exc

    try:
        presentation = Presentation(source_path)
        pages: list[ExtractedPage] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            blocks: list[str] = []
            for shape in sorted(slide.shapes, key=_shape_position):
                blocks.extend(_pptx_shape_text_blocks(shape))
            pages.append(
                ExtractedPage(
                    page_number=slide_number,
                    text="\n".join(blocks),
                )
            )
        return pages
    except Exception as exc:
        raise RuntimeError(
            f"python-pptx failed to extract text from {source_path.name}: {exc}"
        ) from exc


def extract_pptx_text(presentation_path: str | Path) -> ExtractedDocument:
    """Extract PPTX text while preserving slide numbers."""
    source_path = Path(presentation_path).expanduser().resolve()
    if not source_path.exists():
        raise FileNotFoundError(f"PowerPoint presentation not found: {source_path}")
    if not source_path.is_file():
        raise ValueError(f"Expected a file path, got: {source_path}")

    if source_path.suffix.lower() != ".pptx":
        raise ValueError(f"Expected a PPTX file, got: {source_path.name}")

    pages = _extract_pptx_pages(source_path)
    if not pages or not any(page.text for page in pages):
        raise RuntimeError(
            f"PowerPoint extraction produced no text from {source_path.name}"
        )

    return ExtractedDocument(
        source_path=str(source_path),
        page_count=len(pages),
        pages=pages,
    )
